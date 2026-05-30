from __future__ import annotations

import csv
import io
import json
import os
import sqlite3
from datetime import date, datetime, timedelta, timezone
from functools import wraps
from typing import Any

from flask import (
    Flask,
    Response,
    flash,
    jsonify,
    redirect,
    render_template,
    request,
    send_file,
    session,
    url_for,
)

try:
    from openpyxl import Workbook, load_workbook
except ImportError:  # pragma: no cover - optional until requirements are installed
    Workbook = None
    load_workbook = None


DATABASE = os.environ.get("LIFECYCLE_DB", os.path.join(os.path.dirname(__file__), "lifecycle_tracker.db"))
DATE_FIELDS = {"start_date", "due_date", "completed_date"}
REQUIRED_FIELDS = {"tracker_type", "title", "stage", "status", "owner"}
TRACKER_TYPES = {"I20", "D20"}
ROLES = {"Admin": "admin", "Editor": "editor", "Viewer": "viewer"}

FIELD_ALIASES = {
    "tracker type": "tracker_type",
    "trackertype": "tracker_type",
    "tracker": "tracker_type",
    "type": "tracker_type",
    "activity": "activity",
    "opportunity name": "title",
    "opportunity": "title",
    "title": "title",
    "name": "title",
    "account name": "customer",
    "customer": "customer",
    "customer name": "customer",
    "account": "customer",
    "owner": "owner",
    "assigned to": "owner",
    "cortave owner": "cortave_owner",
    "innovator owner": "innovator_owner",
    "who @ cortave": "who_at_cortave",
    "who cortave": "who_at_cortave",
    "stage": "stage",
    "status": "status",
    "description": "description",
    "notes": "notes",
    "start day": "start_day",
    "start date": "start_date",
    "startdate": "start_date",
    "actual start date": "actual_start_date",
    "actual start date edit kick off date for automated deadlines": "actual_start_date",
    "due date": "due_date",
    "duedate": "due_date",
    "target date": "due_date",
    "target due day": "target_due_day",
    "completed date": "completed_date",
    "completion date": "completed_date",
    "completeddate": "completed_date",
    "priority": "priority",
    "next action": "next_action",
    "nextaction": "next_action",
    "link": "link",
    "url": "link",
}

RECORD_FIELDS = [
    "tracker_type",
    "stage",
    "status",
    "title",
    "activity",
    "owner",
    "who_at_cortave",
    "cortave_owner",
    "innovator_owner",
    "customer",
    "description",
    "notes",
    "start_day",
    "start_date",
    "actual_start_date",
    "target_due_day",
    "due_date",
    "completed_date",
    "priority",
    "next_action",
    "link",
    "additional_fields",
]


def create_app(test_config: dict[str, Any] | None = None) -> Flask:
    app = Flask(__name__)
    app.config.update(
        SECRET_KEY=os.environ.get("SECRET_KEY", "dev-lifecycle-tracker-key"),
        DATABASE=DATABASE,
        MAX_CONTENT_LENGTH=16 * 1024 * 1024,
    )
    if test_config:
        app.config.update(test_config)

    @app.before_request
    def ensure_database_and_default_user() -> None:
        init_db(app.config["DATABASE"])
        if "user" not in session:
            session["user"] = {"name": "Lifecycle Admin", "role": "admin"}

    @app.context_processor
    def inject_user() -> dict[str, Any]:
        return {"current_user": session.get("user", {"name": "Lifecycle Admin", "role": "admin"})}

    @app.route("/")
    def dashboard() -> str:
        return render_template("dashboard.html", page_title="Overview Dashboard")

    @app.route("/trackers/<tracker_type>")
    def tracker_page(tracker_type: str) -> str:
        tracker_type = tracker_type.upper()
        if tracker_type not in TRACKER_TYPES:
            return redirect(url_for("dashboard"))
        return render_template("tracker.html", page_title=f"{tracker_type} Tracker", tracker_type=tracker_type)

    @app.route("/settings")
    @role_required("admin")
    def settings_page() -> str:
        return render_template("settings.html", page_title="Data Validation / Settings")

    @app.route("/import-export")
    @role_required("admin")
    def import_export_page() -> str:
        return render_template("import_export.html", page_title="Import / Export")

    @app.route("/admin")
    @role_required("admin")
    def admin_page() -> str:
        return render_template("admin.html", page_title="Admin Settings", roles=ROLES)

    @app.route("/login", methods=["GET", "POST"])
    def login() -> Response | str:
        if request.method == "POST":
            role = request.form.get("role", "admin")
            if role not in ROLES.values():
                flash("Invalid role selected.", "error")
                return redirect(url_for("login"))
            name = request.form.get("name", "").strip() or role.title()
            session["user"] = {"name": name, "role": role}
            return redirect(url_for("dashboard"))
        return render_template("login.html", page_title="Sign in", roles=ROLES)

    @app.route("/logout", methods=["POST"])
    def logout() -> Response:
        session.clear()
        return redirect(url_for("login"))

    @app.get("/api/options")
    def get_options() -> Response:
        db = get_db(app.config["DATABASE"])
        rows = db.execute(
            """
            SELECT id, category, value, label, color, sort_order
            FROM settings_options
            WHERE is_active = 1
            ORDER BY category, sort_order, label
            """
        ).fetchall()
        grouped: dict[str, list[dict[str, Any]]] = {}
        for row in rows:
            grouped.setdefault(row["category"], []).append(dict(row))
        return jsonify(grouped)

    @app.post("/api/options")
    @role_required("admin")
    def create_option() -> Response:
        payload = request.get_json(force=True)
        category = clean_text(payload.get("category"))
        value = clean_text(payload.get("value"))
        label = clean_text(payload.get("label")) or value
        color = clean_text(payload.get("color"))
        sort_order = int(payload.get("sortOrder") or payload.get("sort_order") or 100)
        if category not in {"status", "stage", "owner", "priority"} or not value:
            return jsonify({"error": "Category and value are required."}), 400
        db = get_db(app.config["DATABASE"])
        try:
            cursor = db.execute(
                """
                INSERT INTO settings_options (category, value, label, color, sort_order, is_active)
                VALUES (?, ?, ?, ?, ?, 1)
                """,
                (category, value, label, color, sort_order),
            )
            db.commit()
        except sqlite3.IntegrityError:
            return jsonify({"error": "That option already exists."}), 409
        return jsonify({"id": cursor.lastrowid}), 201

    @app.put("/api/options/<int:option_id>")
    @role_required("admin")
    def update_option(option_id: int) -> Response:
        payload = request.get_json(force=True)
        value = clean_text(payload.get("value"))
        label = clean_text(payload.get("label")) or value
        color = clean_text(payload.get("color"))
        sort_order = int(payload.get("sortOrder") or payload.get("sort_order") or 100)
        if not value:
            return jsonify({"error": "Value is required."}), 400
        db = get_db(app.config["DATABASE"])
        db.execute(
            """
            UPDATE settings_options
            SET value = ?, label = ?, color = ?, sort_order = ?
            WHERE id = ?
            """,
            (value, label, color, sort_order, option_id),
        )
        db.commit()
        return jsonify({"ok": True})

    @app.delete("/api/options/<int:option_id>")
    @role_required("admin")
    def delete_option(option_id: int) -> Response:
        db = get_db(app.config["DATABASE"])
        db.execute("UPDATE settings_options SET is_active = 0 WHERE id = ?", (option_id,))
        db.commit()
        return jsonify({"ok": True})

    @app.get("/api/dashboard")
    def dashboard_data() -> Response:
        db = get_db(app.config["DATABASE"])
        today = date.today().isoformat()
        week_end = (date.today() + timedelta(days=7)).isoformat()
        records = [record_to_dict(row) for row in db.execute("SELECT * FROM tracker_records ORDER BY updated_at DESC").fetchall()]
        summary = {
            "totalI20": count_where(records, trackerType="I20"),
            "totalD20": count_where(records, trackerType="D20"),
            "byStatus": group_count(records, "status"),
            "byStage": group_count(records, "stage"),
            "overdue": [
                record for record in records if record.get("dueDate") and record["dueDate"] < today and record["status"] != "Completed"
            ],
            "dueThisWeek": [
                record
                for record in records
                if record.get("dueDate") and today <= record["dueDate"] <= week_end and record["status"] != "Completed"
            ],
            "recentlyUpdated": records[:8],
            "completed": count_where(records, status="Completed"),
            "onHold": count_where(records, status="On Hold"),
            "inProgress": count_where(records, status="In Progress"),
            "notStarted": count_where(records, status="Not Started"),
        }
        return jsonify(summary)

    @app.get("/api/records")
    def list_records() -> Response:
        filters = parse_record_filters(request.args)
        records = query_records(app.config["DATABASE"], filters)
        return jsonify({"records": records})

    @app.get("/api/records/<int:record_id>")
    def get_record(record_id: int) -> Response:
        db = get_db(app.config["DATABASE"])
        row = db.execute("SELECT * FROM tracker_records WHERE id = ?", (record_id,)).fetchone()
        if not row:
            return jsonify({"error": "Record not found."}), 404
        history = [
            dict(item)
            for item in db.execute(
                "SELECT * FROM audit_history WHERE tracker_record_id = ? ORDER BY changed_at DESC",
                (record_id,),
            ).fetchall()
        ]
        return jsonify({"record": record_to_dict(row), "history": history})

    @app.post("/api/records")
    @role_required("admin", "editor")
    def create_record() -> Response:
        payload = request.get_json(force=True)
        record, errors = validate_record_payload(payload, app.config["DATABASE"])
        if errors:
            return jsonify({"errors": errors}), 400
        now = utc_now()
        user = session["user"]["name"]
        db = get_db(app.config["DATABASE"])
        values = [record.get(field) for field in RECORD_FIELDS]
        cursor = db.execute(
            f"""
            INSERT INTO tracker_records ({", ".join(RECORD_FIELDS)}, created_at, updated_at, updated_by)
            VALUES ({", ".join(["?"] * len(RECORD_FIELDS))}, ?, ?, ?)
            """,
            (*values, now, now, user),
        )
        record_id = cursor.lastrowid
        db.execute(
            """
            INSERT INTO audit_history
                (tracker_record_id, changed_at, changed_by, previous_status, new_status, previous_stage, new_stage, change_summary)
            VALUES (?, ?, ?, NULL, ?, NULL, ?, ?)
            """,
            (record_id, now, user, record["status"], record["stage"], json.dumps({"action": "created"})),
        )
        db.commit()
        return jsonify({"id": record_id}), 201

    @app.put("/api/records/<int:record_id>")
    @role_required("admin", "editor")
    def update_record(record_id: int) -> Response:
        payload = request.get_json(force=True)
        record, errors = validate_record_payload(payload, app.config["DATABASE"])
        if errors:
            return jsonify({"errors": errors}), 400
        db = get_db(app.config["DATABASE"])
        existing = db.execute("SELECT * FROM tracker_records WHERE id = ?", (record_id,)).fetchone()
        if not existing:
            return jsonify({"error": "Record not found."}), 404
        now = utc_now()
        user = session["user"]["name"]
        db.execute(
            f"""
            UPDATE tracker_records
            SET {", ".join([field + " = ?" for field in RECORD_FIELDS])},
                updated_at = ?,
                updated_by = ?
            WHERE id = ?
            """,
            (*[record.get(field) for field in RECORD_FIELDS], now, user, record_id),
        )
        changes = summarize_changes(existing, record)
        if changes:
            db.execute(
                """
                INSERT INTO audit_history
                    (tracker_record_id, changed_at, changed_by, previous_status, new_status, previous_stage, new_stage, change_summary)
                VALUES (?, ?, ?, ?, ?, ?, ?, ?)
                """,
                (
                    record_id,
                    now,
                    user,
                    existing["status"],
                    record["status"],
                    existing["stage"],
                    record["stage"],
                    json.dumps(changes),
                ),
            )
        db.commit()
        return jsonify({"ok": True})

    @app.delete("/api/records/<int:record_id>")
    @role_required("admin")
    def delete_record(record_id: int) -> Response:
        db = get_db(app.config["DATABASE"])
        db.execute("DELETE FROM tracker_records WHERE id = ?", (record_id,))
        db.commit()
        return jsonify({"ok": True})

    @app.post("/api/import")
    @role_required("admin")
    def import_records() -> Response:
        upload = request.files.get("file")
        tracker_type = clean_text(request.form.get("trackerType") or request.form.get("tracker_type")).upper()
        if tracker_type not in TRACKER_TYPES:
            return jsonify({"error": "Choose I20 or D20 as the import target."}), 400
        if not upload or not upload.filename:
            return jsonify({"error": "Upload a CSV or Excel file."}), 400

        rows = parse_uploaded_file(upload, tracker_type)
        imported = 0
        errors = []
        for index, row in enumerate(rows, start=2):
            payload = map_import_row(row, tracker_type)
            record, row_errors = validate_record_payload(payload, app.config["DATABASE"])
            if row_errors:
                errors.append({"row": index, "errors": row_errors})
                continue
            create_record_from_import(record, app.config["DATABASE"], session["user"]["name"])
            imported += 1
        return jsonify({"imported": imported, "errors": errors})

    @app.get("/api/export")
    @role_required("admin", "editor", "viewer")
    def export_records() -> Response:
        filters = parse_record_filters(request.args)
        records = query_records(app.config["DATABASE"], filters)
        export_format = request.args.get("format", "csv").lower()
        filename_prefix = (filters.get("tracker_type") or "lifecycle").lower()
        flat_records = flatten_export_records(records)
        if export_format == "xlsx":
            if Workbook is None:
                return jsonify({"error": "Excel export requires openpyxl."}), 500
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "Tracker Records"
            headers = export_headers(records)
            sheet.append(headers)
            for record in flat_records:
                sheet.append([record.get(header) for header in headers])
            output = io.BytesIO()
            workbook.save(output)
            output.seek(0)
            return send_file(
                output,
                as_attachment=True,
                download_name=f"{filename_prefix}-tracker-export.xlsx",
                mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )

        output = io.StringIO()
        headers = export_headers(records)
        writer = csv.DictWriter(output, fieldnames=headers, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(flat_records)
        return Response(
            output.getvalue(),
            mimetype="text/csv",
            headers={"Content-Disposition": f"attachment; filename={filename_prefix}-tracker-export.csv"},
        )

    return app


def role_required(*roles: str):
    def decorator(view):
        @wraps(view)
        def wrapped(*args, **kwargs):
            user = session.get("user", {"role": "admin"})
            if user.get("role") not in roles:
                if request.path.startswith("/api/"):
                    return jsonify({"error": "You do not have permission for this action."}), 403
                flash("You do not have permission for that page.", "error")
                return redirect(url_for("dashboard"))
            return view(*args, **kwargs)

        return wrapped

    return decorator


def get_db(database_path: str) -> sqlite3.Connection:
    connection = sqlite3.connect(database_path)
    connection.row_factory = sqlite3.Row
    return connection


def init_db(database_path: str) -> None:
    db = get_db(database_path)
    db.execute(
        """
        CREATE TABLE IF NOT EXISTS tracker_records (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            tracker_type TEXT NOT NULL,
            stage TEXT NOT NULL,
            status TEXT NOT NULL,
            title TEXT NOT NULL,
            activity TEXT,
            owner TEXT NOT NULL,
            who_at_cortave TEXT,
            cortave_owner TEXT,
            innovator_owner TEXT,
            customer TEXT,
            description TEXT,
            notes TEXT,
            start_day TEXT,
            start_date TEXT,
            actual_start_date TEXT,
            target_due_day TEXT,
            due_date TEXT,
            completed_date TEXT,
            priority TEXT,
            next_action TEXT,
            link TEXT,
            additional_fields TEXT,
            created_at TEXT NOT NULL,
            updated_at TEXT NOT NULL,
            updated_by TEXT
        )
        """
    )
    db.execute(
        """
        CREATE TABLE IF NOT EXISTS settings_options (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            category TEXT NOT NULL,
            value TEXT NOT NULL,
            label TEXT NOT NULL,
            color TEXT,
            sort_order INTEGER NOT NULL DEFAULT 100,
            is_active INTEGER NOT NULL DEFAULT 1,
            UNIQUE(category, value)
        )
        """
    )
    db.execute(
        """
        CREATE TABLE IF NOT EXISTS audit_history (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            tracker_record_id INTEGER NOT NULL,
            changed_at TEXT NOT NULL,
            changed_by TEXT,
            previous_status TEXT,
            new_status TEXT,
            previous_stage TEXT,
            new_stage TEXT,
            change_summary TEXT,
            FOREIGN KEY(tracker_record_id) REFERENCES tracker_records(id) ON DELETE CASCADE
        )
        """
    )
    ensure_tracker_record_columns(db)
    seed_options(db)
    db.commit()
    db.close()


def ensure_tracker_record_columns(db: sqlite3.Connection) -> None:
    existing = {row[1] for row in db.execute("PRAGMA table_info(tracker_records)").fetchall()}
    columns = {
        "activity": "TEXT",
        "who_at_cortave": "TEXT",
        "cortave_owner": "TEXT",
        "innovator_owner": "TEXT",
        "start_day": "TEXT",
        "actual_start_date": "TEXT",
        "target_due_day": "TEXT",
        "link": "TEXT",
    }
    for name, column_type in columns.items():
        if name not in existing:
            db.execute(f"ALTER TABLE tracker_records ADD COLUMN {name} {column_type}")


def seed_options(db: sqlite3.Connection) -> None:
    defaults = [
        ("status", "Not Started", "Not Started", "#f4c7c3", 10),
        ("status", "In Progress", "In Progress", "#fce8b2", 20),
        ("status", "On Hold", "On Hold", "#c9daf8", 30),
        ("status", "Completed", "Completed", "#b7e1cd", 40),
        ("stage", "I20", "I20", "", 10),
        ("stage", "I50", "I50", "", 20),
        ("stage", "D20", "D20", "", 30),
        ("stage", "Qualified Out", "Qualified Out", "", 40),
        ("stage", "I0", "I0", "", 50),
        ("stage", "D0", "D0", "", 60),
        ("stage", "D50", "D50", "", 70),
        ("owner", "Unassigned", "Unassigned", "", 10),
        ("owner", "Wade / Sales", "Wade / Sales", "", 20),
        ("owner", "Melissa / Marketing", "Melissa / Marketing", "", 30),
        ("owner", "Mark / Innovator", "Mark / Innovator", "", 40),
        ("owner", "Bobby / Experience", "Bobby / Experience", "", 50),
        ("owner", "Nick / Sales", "Nick / Sales", "", 60),
        ("owner", "Customer", "Customer", "", 70),
        ("owner", "Legal", "Legal", "", 80),
        ("owner", "Finance", "Finance", "", 90),
        ("owner", "Innovator", "Innovator", "", 100),
        ("owner", "Innovator / cortave", "Innovator / cortave", "", 110),
        ("priority", "Low", "Low", "", 10),
        ("priority", "Medium", "Medium", "", 20),
        ("priority", "High", "High", "", 30),
        ("priority", "Critical", "Critical", "", 40),
    ]
    db.executemany(
        """
        INSERT OR IGNORE INTO settings_options (category, value, label, color, sort_order)
        VALUES (?, ?, ?, ?, ?)
        """,
        defaults,
    )


def clean_text(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()


def utc_now() -> str:
    return datetime.now(timezone.utc).replace(microsecond=0).isoformat().replace("+00:00", "Z")


def validate_record_payload(payload: dict[str, Any], database_path: str) -> tuple[dict[str, Any], list[str]]:
    raw_dates = {
        "start_date": payload.get("startDate") or payload.get("start_date"),
        "actual_start_date": payload.get("actualStartDate") or payload.get("actual_start_date"),
        "due_date": payload.get("dueDate") or payload.get("due_date"),
        "completed_date": payload.get("completedDate") or payload.get("completed_date"),
    }
    record = {
        "tracker_type": clean_text(payload.get("trackerType") or payload.get("tracker_type")).upper(),
        "stage": clean_text(payload.get("stage")),
        "status": clean_text(payload.get("status")),
        "title": clean_text(payload.get("title") or payload.get("opportunityName") or payload.get("activity")),
        "activity": clean_text(payload.get("activity") or payload.get("title") or payload.get("opportunityName")),
        "owner": clean_text(payload.get("owner") or payload.get("cortaveOwner") or payload.get("cortave_owner") or payload.get("whoAtCortave") or payload.get("who_at_cortave")),
        "who_at_cortave": clean_text(payload.get("whoAtCortave") or payload.get("who_at_cortave")),
        "cortave_owner": clean_text(payload.get("cortaveOwner") or payload.get("cortave_owner")),
        "innovator_owner": clean_text(payload.get("innovatorOwner") or payload.get("innovator_owner")),
        "customer": clean_text(payload.get("customer") or payload.get("accountName")),
        "description": clean_text(payload.get("description")),
        "notes": clean_text(payload.get("notes")),
        "start_day": clean_text(payload.get("startDay") or payload.get("start_day")),
        "start_date": clean_date(raw_dates["start_date"]),
        "actual_start_date": clean_date(raw_dates["actual_start_date"]),
        "target_due_day": clean_text(payload.get("targetDueDay") or payload.get("target_due_day")),
        "due_date": clean_date(raw_dates["due_date"]),
        "completed_date": clean_date(raw_dates["completed_date"]),
        "priority": clean_text(payload.get("priority")),
        "next_action": clean_text(payload.get("nextAction") or payload.get("next_action")),
        "link": clean_text(payload.get("link")),
        "additional_fields": normalize_additional_fields(payload.get("additionalFields") or payload.get("additional_fields")),
    }
    errors = []
    for field in sorted(REQUIRED_FIELDS):
        if not record[field]:
            errors.append(f"{field.replace('_', ' ').title()} is required.")
    for field, raw_value in raw_dates.items():
        if clean_text(raw_value) and not record[field]:
            errors.append(f"{field.replace('_', ' ').title()} must be a valid date.")
    if record["tracker_type"] and record["tracker_type"] not in TRACKER_TYPES:
        errors.append("Tracker type must be I20 or D20.")
    option_values = load_option_values(database_path)
    if record["status"] and record["status"] not in option_values["status"]:
        errors.append("Status must be one of the configured status options.")
    if record["stage"] and record["stage"] not in option_values["stage"]:
        errors.append("Stage must be one of the configured stage options.")
    if record["priority"] and record["priority"] not in option_values["priority"]:
        errors.append("Priority must be one of the configured priority options.")
    if record["status"] == "Completed" and not record["completed_date"]:
        record["completed_date"] = date.today().isoformat()
    return record, errors


def clean_date(value: Any) -> str:
    text = clean_text(value)
    if not text:
        return ""
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    for fmt in ("%Y-%m-%d", "%m/%d/%Y", "%d/%m/%Y", "%m-%d-%Y", "%d-%m-%Y"):
        try:
            return datetime.strptime(text, fmt).date().isoformat()
        except ValueError:
            continue
    return ""


def normalize_additional_fields(value: Any) -> str:
    if isinstance(value, dict):
        return json.dumps(value, sort_keys=True)
    text = clean_text(value)
    if not text:
        return "{}"
    try:
        parsed = json.loads(text)
        if isinstance(parsed, dict):
            return json.dumps(parsed, sort_keys=True)
    except json.JSONDecodeError:
        pass
    return json.dumps({"notes": text})


def load_option_values(database_path: str) -> dict[str, set[str]]:
    db = get_db(database_path)
    rows = db.execute("SELECT category, value FROM settings_options WHERE is_active = 1").fetchall()
    values = {"status": set(), "stage": set(), "owner": set(), "priority": set()}
    for row in rows:
        values.setdefault(row["category"], set()).add(row["value"])
    db.close()
    return values


def parse_record_filters(args: Any) -> dict[str, Any]:
    return {
        "tracker_type": clean_text(args.get("trackerType") or args.get("tracker_type")).upper(),
        "status": clean_text(args.get("status")),
        "stage": clean_text(args.get("stage")),
        "owner": clean_text(args.get("owner")),
        "search": clean_text(args.get("search")),
        "due_before": clean_date(args.get("dueBefore") or args.get("due_before")),
        "due_after": clean_date(args.get("dueAfter") or args.get("due_after")),
        "overdue": clean_text(args.get("overdue")).lower() in {"true", "1", "yes"},
        "sort": clean_text(args.get("sort")) or "updated_at",
        "direction": "ASC" if clean_text(args.get("direction")).lower() == "asc" else "DESC",
    }


def query_records(database_path: str, filters: dict[str, Any]) -> list[dict[str, Any]]:
    clauses = []
    params: list[Any] = []
    if filters.get("tracker_type") in TRACKER_TYPES:
        clauses.append("tracker_type = ?")
        params.append(filters["tracker_type"])
    for key in ("status", "stage", "owner"):
        if filters.get(key):
            clauses.append(f"{key} = ?")
            params.append(filters[key])
    if filters.get("search"):
        clauses.append("(title LIKE ? OR customer LIKE ? OR owner LIKE ? OR description LIKE ? OR notes LIKE ?)")
        search = f"%{filters['search']}%"
        params.extend([search, search, search, search, search])
    if filters.get("due_before"):
        clauses.append("due_date <= ?")
        params.append(filters["due_before"])
    if filters.get("due_after"):
        clauses.append("due_date >= ?")
        params.append(filters["due_after"])
    if filters.get("overdue"):
        clauses.append("due_date < ? AND status != 'Completed'")
        params.append(date.today().isoformat())
    where = f"WHERE {' AND '.join(clauses)}" if clauses else ""
    sort_map = {
        "title": "title",
        "stage": "stage",
        "status": "status",
        "owner": "owner",
        "customer": "customer",
        "dueDate": "due_date",
        "due_date": "due_date",
        "updatedAt": "updated_at",
        "updated_at": "updated_at",
        "priority": "priority",
    }
    sort_column = sort_map.get(filters.get("sort"), "updated_at")
    db = get_db(database_path)
    rows = db.execute(
        f"SELECT * FROM tracker_records {where} ORDER BY {sort_column} {filters['direction']}",
        params,
    ).fetchall()
    db.close()
    return [record_to_dict(row) for row in rows]


def record_to_dict(row: sqlite3.Row) -> dict[str, Any]:
    try:
        additional = json.loads(row["additional_fields"] or "{}")
    except json.JSONDecodeError:
        additional = {}
    return {
        "id": row["id"],
        "trackerType": row["tracker_type"],
        "stage": row["stage"],
        "status": row["status"],
        "title": row["title"],
        "activity": row["activity"] or row["title"],
        "owner": row["owner"],
        "whoAtCortave": row["who_at_cortave"] or "",
        "cortaveOwner": row["cortave_owner"] or "",
        "innovatorOwner": row["innovator_owner"] or "",
        "customer": row["customer"] or "",
        "description": row["description"] or "",
        "notes": row["notes"] or "",
        "startDay": row["start_day"] or "",
        "startDate": row["start_date"] or "",
        "actualStartDate": row["actual_start_date"] or "",
        "targetDueDay": row["target_due_day"] or "",
        "dueDate": row["due_date"] or "",
        "completedDate": row["completed_date"] or "",
        "priority": row["priority"] or "",
        "nextAction": row["next_action"] or "",
        "link": row["link"] or "",
        "additionalFields": additional,
        "createdAt": row["created_at"],
        "updatedAt": row["updated_at"],
        "updatedBy": row["updated_by"] or "",
        "isOverdue": bool(row["due_date"] and row["due_date"] < date.today().isoformat() and row["status"] != "Completed"),
    }


def count_where(records: list[dict[str, Any]], **criteria: Any) -> int:
    return sum(1 for record in records if all(record.get(key) == value for key, value in criteria.items()))


def group_count(records: list[dict[str, Any]], field: str) -> dict[str, int]:
    counts: dict[str, int] = {}
    for record in records:
        counts[record.get(field) or "Unspecified"] = counts.get(record.get(field) or "Unspecified", 0) + 1
    return counts


def summarize_changes(existing: sqlite3.Row, updated: dict[str, Any]) -> dict[str, Any]:
    changes = {}
    for field in RECORD_FIELDS:
        if clean_text(existing[field]) != clean_text(updated.get(field)):
            changes[field] = {"previous": existing[field], "new": updated.get(field)}
    return changes


def parse_uploaded_file(upload: Any, tracker_type: str = "") -> list[dict[str, Any]]:
    filename = upload.filename.lower()
    if filename.endswith(".xlsx"):
        if load_workbook is None:
            raise ValueError("Excel import requires openpyxl.")
        workbook = load_workbook(upload.stream, data_only=True)
        preferred_sheet = f"{tracker_type.upper()} Tracker" if tracker_type else ""
        sheet = workbook[preferred_sheet] if preferred_sheet in workbook.sheetnames else workbook.active
        rows = list(sheet.iter_rows(values_only=True))
        if not rows:
            return []
        header_index = find_header_row(rows)
        headers = [clean_text(value) for value in rows[header_index]]
        return [dict(zip(headers, row)) for row in rows[header_index + 1 :] if any(cell not in (None, "") for cell in row)]
    stream = io.StringIO(upload.stream.read().decode("utf-8-sig"))
    return list(csv.DictReader(stream))


def find_header_row(rows: list[tuple[Any, ...]]) -> int:
    for index, row in enumerate(rows[:20]):
        normalized = {normalize_header(value) for value in row if clean_text(value)}
        if {"stage", "status"}.issubset(normalized) and ("activity" in normalized or "opportunity name" in normalized or "title" in normalized):
            return index
    return 0


def map_import_row(row: dict[str, Any], tracker_type: str) -> dict[str, Any]:
    payload: dict[str, Any] = {"trackerType": tracker_type}
    extra: dict[str, Any] = {}
    for raw_key, value in row.items():
        key = clean_text(raw_key)
        mapped = FIELD_ALIASES.get(normalize_header(key))
        if mapped:
            payload[to_camel(mapped)] = value
        elif key:
            extra[key] = value
    payload["trackerType"] = tracker_type
    payload["additionalFields"] = extra
    return payload


def normalize_header(value: Any) -> str:
    text = clean_text(value).lower().replace("_", " ").replace("@", " @ ")
    normalized = "".join(char if char.isalnum() or char in {" ", "@"} else " " for char in text)
    return " ".join(normalized.split())


def to_camel(field: str) -> str:
    parts = field.split("_")
    return parts[0] + "".join(part.title() for part in parts[1:])


def create_record_from_import(record: dict[str, Any], database_path: str, user: str) -> None:
    now = utc_now()
    db = get_db(database_path)
    values = [record.get(field) for field in RECORD_FIELDS]
    cursor = db.execute(
        f"""
        INSERT INTO tracker_records ({", ".join(RECORD_FIELDS)}, created_at, updated_at, updated_by)
        VALUES ({", ".join(["?"] * len(RECORD_FIELDS))}, ?, ?, ?)
        """,
        (*values, now, now, user),
    )
    db.execute(
        """
        INSERT INTO audit_history
            (tracker_record_id, changed_at, changed_by, previous_status, new_status, previous_stage, new_stage, change_summary)
        VALUES (?, ?, ?, NULL, ?, NULL, ?, ?)
        """,
        (cursor.lastrowid, now, user, record["status"], record["stage"], json.dumps({"action": "imported"})),
    )
    db.commit()
    db.close()


def export_headers(records: list[dict[str, Any]]) -> list[str]:
    headers = [
        "id",
        "trackerType",
        "stage",
        "status",
        "title",
        "activity",
        "owner",
        "whoAtCortave",
        "cortaveOwner",
        "innovatorOwner",
        "customer",
        "description",
        "notes",
        "startDay",
        "startDate",
        "actualStartDate",
        "targetDueDay",
        "dueDate",
        "completedDate",
        "priority",
        "nextAction",
        "link",
        "createdAt",
        "updatedAt",
        "updatedBy",
    ]
    extra_keys = sorted({key for record in records for key in record.get("additionalFields", {}).keys()})
    return headers + extra_keys


def flatten_export_records(records: list[dict[str, Any]]) -> list[dict[str, Any]]:
    flattened = []
    for record in records:
        row = {key: value for key, value in record.items() if key != "additionalFields"}
        row.update(record.get("additionalFields", {}))
        flattened.append(row)
    return flattened


if __name__ == "__main__":
    app = create_app()
    app.run(host="0.0.0.0", port=int(os.environ.get("PORT", 5000)), debug=os.environ.get("FLASK_DEBUG") == "1")
