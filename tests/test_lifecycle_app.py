import io
import os
from pathlib import Path
import tempfile
import unittest

from openpyxl import Workbook

from app import create_app, init_db


class LifecycleAppTest(unittest.TestCase):
    def setUp(self):
        self.tempdir = tempfile.TemporaryDirectory()
        self.database = os.path.join(self.tempdir.name, "test.db")
        init_db(self.database)
        self.app = create_app({"TESTING": True, "DATABASE": self.database, "SECRET_KEY": "test"})
        self.client = self.app.test_client()

    def tearDown(self):
        self.tempdir.cleanup()

    def test_seeded_options_include_required_statuses_and_stages(self):
        response = self.client.get("/api/options")
        self.assertEqual(response.status_code, 200)
        payload = response.get_json()
        statuses = {item["value"] for item in payload["status"]}
        stages = {item["value"] for item in payload["stage"]}
        self.assertTrue({"Not Started", "In Progress", "On Hold", "Completed"}.issubset(statuses))
        self.assertTrue({"I20", "I50", "D20", "Qualified Out", "I0", "D0", "D50"}.issubset(stages))

    def test_record_crud_and_validation(self):
        bad_response = self.client.post(
            "/api/records",
            json={"trackerType": "I20", "title": "Bad date", "stage": "I20", "status": "Not Started", "owner": "Owner", "dueDate": "bad"},
        )
        self.assertEqual(bad_response.status_code, 400)
        self.assertIn("valid date", " ".join(bad_response.get_json()["errors"]))

        create_response = self.client.post(
            "/api/records",
            json={
                "trackerType": "I20",
                "title": "Lifecycle opportunity",
                "stage": "I20",
                "status": "In Progress",
                "owner": "Owner",
                "customer": "Customer",
                "dueDate": "2030-01-15",
                "priority": "High",
            },
        )
        self.assertEqual(create_response.status_code, 201)
        record_id = create_response.get_json()["id"]

        list_response = self.client.get("/api/records?trackerType=I20&status=In%20Progress")
        records = list_response.get_json()["records"]
        self.assertEqual(len(records), 1)
        self.assertEqual(records[0]["title"], "Lifecycle opportunity")

        update_response = self.client.put(
            f"/api/records/{record_id}",
            json={
                "trackerType": "I20",
                "title": "Lifecycle opportunity",
                "stage": "I50",
                "status": "Completed",
                "owner": "Owner",
                "customer": "Customer",
                "dueDate": "2030-01-15",
                "priority": "High",
            },
        )
        self.assertEqual(update_response.status_code, 200)

        detail_response = self.client.get(f"/api/records/{record_id}")
        detail = detail_response.get_json()
        self.assertEqual(detail["record"]["status"], "Completed")
        self.assertTrue(detail["record"]["completedDate"])
        self.assertGreaterEqual(len(detail["history"]), 1)

    def test_csv_import_preserves_unknown_columns(self):
        csv_bytes = b"Opportunity Name,Owner,Stage,Status,Due Date,Workbook Extra\nImported item,Owner,I20,Not Started,2030-02-01,Preserved\n"
        response = self.client.post(
            "/api/import",
            data={"trackerType": "I20", "file": (io.BytesIO(csv_bytes), "records.csv")},
            content_type="multipart/form-data",
        )
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.get_json()["imported"], 1)

        records = self.client.get("/api/records?trackerType=I20").get_json()["records"]
        self.assertEqual(records[0]["additionalFields"]["Workbook Extra"], "Preserved")

    def test_lifecycle_workbook_import_uses_tracker_sheet_and_fields(self):
        workbook = Workbook()
        overview = workbook.active
        overview.title = "Overview"
        overview.append(["This sheet should not import"])
        tracker = workbook.create_sheet("D20 Tracker")
        tracker.append([
            "Stage",
            "Activity",
            "Status",
            "who @ cortave ",
            "Start Day",
            "Actual Start Date \n(Edit Kick Off Date for Automated deadlines)",
            "Target Due Day",
            "Due Date",
            "cortave Owner",
            "Innovator Owner",
            "Link",
            "Notes",
        ])
        tracker.append([
            "D20",
            "Discovery meeting",
            "In Progress",
            "Wade / Sales",
            20,
            "2030-01-01",
            10,
            "2030-01-11",
            "Bobby",
            "Innovator contact",
            "https://example.com",
            "Workbook note",
        ])

        org = workbook.create_sheet("Org Impact ")
        org["C4"] = 5
        org["D4"] = "Working days per week"
        org["C10"] = 2008
        org["D10"] = "Working hours per year (average)"
        org["C11"] = 100000
        org["D11"] = "Average FTE cost (annual)"
        org["C16"] = "Partner Manager"
        org["D16"] = 15
        org["E16"] = 45
        org["F16"] = 90
        org["G16"] = 3
        org["H16"] = 9
        org["I16"] = 18
        org["C26"] = "Total"
        org["D26"] = 15
        org["E26"] = 45
        org["F26"] = 90
        org["G26"] = 3
        org["H26"] = 9
        org["I26"] = 18
        org["C31"] = "Partner Manager"
        org["D31"] = 120
        org["E31"] = 360
        org["F31"] = 720
        org["G31"] = 5976.1
        org["H31"] = 17928.3
        org["I31"] = 35856.6
        org["C41"] = "Total"
        org["D41"] = 120
        org["E41"] = 360
        org["F41"] = 720
        org["G41"] = 5976.1
        org["H41"] = 17928.3
        org["I41"] = 35856.6
        org["C45"] = "Recruitment"
        org["D45"] = 40
        org["E45"] = 7360
        org["F45"] = 920
        org["G45"] = 3.67
        org["C49"] = "Total"
        org["D49"] = 49
        org["E49"] = 14856
        org["F49"] = 1857
        org["G49"] = 7.4

        output = io.BytesIO()
        workbook.save(output)
        output.seek(0)

        response = self.client.post(
            "/api/import",
            data={"trackerType": "D20", "file": (output, "Lifecycle Workbook.xlsx")},
            content_type="multipart/form-data",
        )

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.get_json()["imported"], 1)
        records = self.client.get("/api/records?trackerType=D20").get_json()["records"]
        self.assertEqual(records[0]["title"], "Discovery meeting")
        self.assertEqual(records[0]["whoAtCortave"], "Wade / Sales")
        self.assertEqual(records[0]["startDay"], "20")
        self.assertEqual(records[0]["actualStartDate"], "2030-01-01")
        self.assertEqual(records[0]["targetDueDay"], "10")
        self.assertEqual(records[0]["dueDate"], "2030-01-11")
        self.assertEqual(records[0]["cortaveOwner"], "Bobby")
        self.assertEqual(records[0]["innovatorOwner"], "Innovator contact")
        self.assertEqual(records[0]["link"], "https://example.com")

    def test_partner_dashboard_import_upsert_and_linking(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Partners"
        sheet.append([
            "Partner",
            " Owner ",
            "Territory",
            "Active",
            "Date Of Stage Change",
            "Age of Stage",
            "Days Overdue",
            "Workbook Link",
            "SF Account",
            "Next Steps/notes",
            "Partner Type ",
            "CSM Involved ",
        ])
        sheet.append([
            "Deloitte",
            "Mark",
            "EMEA",
            "Onboarding",
            "2023-06-01",
            "#NAME?",
            12,
            "Deloitte Workbook",
            "Existing account",
            "Talk to account team",
            "Martech",
            "No",
        ])
        sheet.append([
            "Deloitte",
            "Melissa",
            "Global",
            "Recruitment",
            "2023-07-01",
            4,
            0,
            "https://example.com/deloitte",
            "Updated account",
            "Updated notes",
            "Technology",
            "Yes",
        ])
        org = workbook.create_sheet("Org Impact ")
        org["C4"] = 5
        org["D4"] = "Working days per week"
        org["C10"] = 2008
        org["D10"] = "Working hours per year (average)"
        org["C11"] = 100000
        org["D11"] = "Average FTE cost (annual)"
        org["C16"] = "Partner Manager"
        org["D16"] = 15
        org["E16"] = 45
        org["F16"] = 90
        org["G16"] = 3
        org["H16"] = 9
        org["I16"] = 18
        org["C26"] = "Total"
        org["D26"] = 15
        org["E26"] = 45
        org["F26"] = 90
        org["G26"] = 3
        org["H26"] = 9
        org["I26"] = 18
        org["C31"] = "Partner Manager"
        org["D31"] = 120
        org["E31"] = 360
        org["F31"] = 720
        org["G31"] = 5976.1
        org["H31"] = 17928.3
        org["I31"] = 35856.6
        org["C41"] = "Total"
        org["D41"] = 120
        org["E41"] = 360
        org["F41"] = 720
        org["G41"] = 5976.1
        org["H41"] = 17928.3
        org["I41"] = 35856.6
        org["C45"] = "Recruitment"
        org["D45"] = 40
        org["E45"] = 7360
        org["F45"] = 920
        org["G45"] = 3.67
        org["C49"] = "Total"
        org["D49"] = 49
        org["E49"] = 14856
        org["F49"] = 1857
        org["G49"] = 7.4

        output = io.BytesIO()
        workbook.save(output)
        output.seek(0)

        response = self.client.post(
            "/api/partners/import",
            data={"file": (output, "Partner Dashboard.xlsx")},
            content_type="multipart/form-data",
        )

        self.assertEqual(response.status_code, 200)
        payload = response.get_json()
        self.assertEqual(payload["created"], 1)
        self.assertEqual(payload["updated"], 1)
        self.assertEqual(payload["errors"], [])

        partners_response = self.client.get("/api/partners?search=Deloitte")
        partners = partners_response.get_json()["partners"]
        self.assertEqual(len(partners), 1)
        self.assertEqual(partners[0]["owner"], "Melissa")
        self.assertEqual(partners[0]["masterStage"], "Recruitment")
        self.assertEqual(partners[0]["workbookLink"], "https://example.com/deloitte")
        self.assertTrue(partners[0]["isWorkbookLinkUrl"])

        org_response = self.client.get("/api/org-impact")
        org_payload = org_response.get_json()
        self.assertEqual(org_payload["totals"]["totalHours"], 14856)
        self.assertEqual(org_payload["summary"][0]["label"], "Recruitment")
        self.assertEqual(org_payload["assumptions"][0]["label"], "Partner Manager")

        link_response = self.client.post("/api/lifecycle-workbook/link", json={"partnerId": partners[0]["id"]})
        self.assertEqual(link_response.status_code, 200)
        workbook_response = self.client.get("/api/lifecycle-workbook")
        linked = workbook_response.get_json()["workbook"]
        self.assertEqual(linked["partner"]["partnerName"], "Deloitte")
        self.assertEqual(linked["partner"]["owner"], "Melissa")

        unlink_response = self.client.post("/api/lifecycle-workbook/unlink", json={})
        self.assertEqual(unlink_response.status_code, 200)
        self.assertIsNone(unlink_response.get_json()["workbook"]["partner"])

    def test_partner_dashboard_requires_partners_sheet(self):
        workbook = Workbook()
        workbook.active.title = "Wrong Sheet"
        output = io.BytesIO()
        workbook.save(output)
        output.seek(0)

        response = self.client.post(
            "/api/partners/import",
            data={"file": (output, "Partner Dashboard.xlsx")},
            content_type="multipart/form-data",
        )

        self.assertEqual(response.status_code, 400)
        self.assertIn("Partners sheet", response.get_json()["error"])

    def test_default_lifecycle_templates_and_account_plan_creation(self):
        templates = self.client.get("/api/lifecycle-templates").get_json()["templates"]
        by_type = {template["accountType"]: template for template in templates}
        self.assertIn("Innovator", by_type)
        self.assertIn("Direct Customer", by_type)

        innovator_items = self.client.get(f"/api/lifecycle-templates/{by_type['Innovator']['id']}/items").get_json()["items"]
        direct_items = self.client.get(f"/api/lifecycle-templates/{by_type['Direct Customer']['id']}/items").get_json()["items"]
        self.assertEqual(len(innovator_items), 27)
        self.assertEqual(len(direct_items), 22)
        self.assertEqual(innovator_items[0]["stage"], "I0")
        self.assertEqual(innovator_items[0]["activity"], "Kick-off date set for automated deadlines")
        self.assertEqual(innovator_items[-1]["activity"], "Innovator Live")
        self.assertEqual(direct_items[0]["stage"], "D0")
        self.assertEqual(direct_items[-1]["activity"], "customer Live")

        create_response = self.client.post(
            "/api/account-plans",
            json={
                "accountName": "Deloitte",
                "accountType": "Innovator",
                "planOwner": "Mark",
                "kickOffDate": "2026-06-01",
                "targetGoLiveDate": "2026-10-01",
                "notes": "Partner plan",
            },
        )
        self.assertEqual(create_response.status_code, 201)
        plan_id = create_response.get_json()["id"]
        items = self.client.get(f"/api/account-plans/{plan_id}/items").get_json()["items"]
        self.assertEqual(len(items), 27)
        self.assertEqual(items[0]["activity"], "Kick-off date set for automated deadlines")
        self.assertEqual(items[0]["actualStartDate"], "2026-06-01")
        self.assertEqual(items[2]["dueDate"], "2026-06-08")
        self.assertEqual(items[-1]["dueDate"], "2026-07-21")

        summary = self.client.get("/api/account-plans").get_json()["summary"]
        self.assertEqual(summary["totalInnovators"], 1)
        self.assertEqual(summary["totalDirectCustomers"], 0)

    def test_direct_customer_account_plan_uses_customer_language(self):
        create_response = self.client.post(
            "/api/account-plans",
            json={
                "accountName": "Acme",
                "accountType": "Direct Customer",
                "planOwner": "Wade",
                "kickOffDate": "2026-06-01",
            },
        )
        self.assertEqual(create_response.status_code, 201)
        items = self.client.get(f"/api/account-plans/{create_response.get_json()['id']}/items").get_json()["items"]
        activities = [item["activity"] for item in items]
        self.assertIn("customer signs EULA agreement", activities)
        self.assertIn("customer Live", activities)
        self.assertNotIn("Innovator signs agreement", activities)

    def test_account_plan_detail_page_and_editing(self):
        create_response = self.client.post(
            "/api/account-plans",
            json={
                "accountName": "Editable Plan",
                "accountType": "Innovator",
                "planOwner": "Mark",
                "kickOffDate": "2026-06-01",
            },
        )
        self.assertEqual(create_response.status_code, 201)
        plan_id = create_response.get_json()["id"]

        page_response = self.client.get(f"/account-plans/{plan_id}")
        self.assertEqual(page_response.status_code, 200)
        self.assertIn(b"Lifecycle Items", page_response.data)

        update_response = self.client.put(
            f"/api/account-plans/{plan_id}",
            json={
                "accountName": "Editable Plan Updated",
                "planOwner": "Melissa",
                "kickOffDate": "2026-06-02",
                "targetGoLiveDate": "2026-08-01",
                "notes": "Updated notes",
            },
        )
        self.assertEqual(update_response.status_code, 200)
        updated_plan = update_response.get_json()["plan"]
        self.assertEqual(updated_plan["accountName"], "Editable Plan Updated")
        self.assertEqual(updated_plan["planOwner"], "Melissa")

        items = self.client.get(f"/api/account-plans/{plan_id}/items").get_json()["items"]
        update_item_response = self.client.put(
            f"/api/lifecycle-items/{items[0]['id']}",
            json={
                "stage": "I0",
                "activity": "Edited lifecycle action",
                "status": "In Progress",
                "responsibleParty": "Mark",
                "actualStartDate": "2026-06-03",
                "dueDate": "2026-06-04",
                "cortaveOwner": "Mark",
                "accountOwner": "Innovator",
                "notes": "Edited item notes",
            },
        )
        self.assertEqual(update_item_response.status_code, 200)
        updated_item = self.client.get(f"/api/account-plans/{plan_id}/items").get_json()["items"][0]
        self.assertEqual(updated_item["activity"], "Edited lifecycle action")
        self.assertEqual(updated_item["status"], "In Progress")
        self.assertEqual(updated_item["dueDate"], "2026-06-04")

    def test_lifecycle_item_status_persists_and_dashboard_recalculates(self):
        first = self.client.post(
            "/api/account-plans",
            json={
                "accountName": "Persistence One",
                "accountType": "Innovator",
                "planOwner": "Mark",
                "kickOffDate": "2020-01-01",
            },
        ).get_json()["plan"]
        second = self.client.post(
            "/api/account-plans",
            json={
                "accountName": "Persistence Two",
                "accountType": "Innovator",
                "planOwner": "Mark",
                "kickOffDate": "2020-01-01",
            },
        ).get_json()["plan"]

        self.assertEqual(len(first["items"]), 27)
        self.assertEqual(len(second["items"]), 27)
        self.assertNotEqual(first["items"][0]["id"], second["items"][0]["id"])

        item_id = first["items"][0]["id"]
        update = self.client.put(
            f"/api/lifecycle-items/{item_id}",
            json={
                "stage": "I0",
                "activity": "Kick-off date set for automated deadlines",
                "status": "In Progress",
                "responsibleParty": "Mark / Innovator",
                "actualStartDate": "2020-01-01",
                "dueDate": "2020-01-01",
                "completedDate": "",
                "notes": "Status changed",
                "link": "",
                "nextAction": "Follow up",
            },
        )
        self.assertEqual(update.status_code, 200)
        self.assertEqual(update.get_json()["item"]["status"], "In Progress")

        refreshed = self.client.get(f"/api/account-plans/{first['id']}/items").get_json()["items"]
        self.assertEqual(refreshed[0]["status"], "In Progress")
        self.assertEqual(refreshed[0]["nextAction"], "Follow up")

        # Saving the plan header must not re-seed or reset lifecycle items.
        self.client.put(
            f"/api/account-plans/{first['id']}",
            json={
                "accountName": "Persistence One Updated",
                "planOwner": "Melissa",
                "kickOffDate": "2020-01-01",
                "targetGoLiveDate": "2020-04-01",
                "territory": "EMEA",
                "notes": "Header update",
            },
        )
        after_header_save = self.client.get(f"/api/account-plans/{first['id']}/items").get_json()["items"]
        self.assertEqual(len(after_header_save), 27)
        self.assertEqual(after_header_save[0]["status"], "In Progress")

        dashboard_plan = next(
            plan for plan in self.client.get("/api/account-plans").get_json()["plans"]
            if plan["id"] == first["id"]
        )
        self.assertEqual(dashboard_plan["healthStatus"], "Overdue")
        self.assertEqual(dashboard_plan["nextStep"], "Kick-off date set for automated deadlines")
        self.assertGreater(dashboard_plan["daysOverdue"], 0)

        completed = self.client.put(
            f"/api/lifecycle-items/{item_id}",
            json={
                "stage": "I0",
                "activity": "Kick-off date set for automated deadlines",
                "status": "Completed",
                "responsibleParty": "Mark / Innovator",
                "actualStartDate": "2020-01-01",
                "dueDate": "2020-01-01",
                "completedDate": "2020-01-02",
                "notes": "Done",
                "link": "",
                "nextAction": "",
            },
        )
        self.assertEqual(completed.get_json()["item"]["status"], "Completed")
        refreshed_completed = self.client.get(f"/api/account-plans/{first['id']}/items").get_json()["items"]
        self.assertEqual(refreshed_completed[0]["status"], "Completed")
        self.assertEqual(refreshed_completed[0]["completedDate"], "2020-01-02")

    def test_status_badge_helper_is_shared_in_frontend(self):
        script = Path("static/app.js").read_text()
        self.assertIn("function getStatusBadgeClass", script)
        self.assertIn("badge-status-in-progress", script)
        self.assertIn("badge-status-completed", script)


if __name__ == "__main__":
    unittest.main()
